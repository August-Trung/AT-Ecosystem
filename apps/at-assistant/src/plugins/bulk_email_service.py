from __future__ import annotations

import json
import os
import re
import unicodedata
from dataclasses import asdict, dataclass
from datetime import datetime
from pathlib import Path
from typing import Any
from urllib.parse import unquote, urlparse

from src.core.app_paths import ensure_runtime_dir


EMAIL_RE = re.compile(r"^[\w\.-]+@[\w\.-]+\.\w+$")


@dataclass
class BulkEmailRow:
    row_number: int
    to: str
    subject: str
    body: str
    candidate_name: str
    greeting_name: str
    cc: str
    bcc: str
    attachment_paths: list[str]


def _normalize_header(value: Any) -> str:
    text = str(value or "").strip().lower()
    normalized = unicodedata.normalize("NFD", text)
    normalized = "".join(ch for ch in normalized if unicodedata.category(ch) != "Mn")
    normalized = re.sub(r"[^a-z0-9]+", " ", normalized)
    return " ".join(normalized.split())


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


class BulkEmailService:
    HEADER_MAP = {
        "tieu de": "subject",
        "ten ung vien": "candidate_name",
        "ten": "candidate_name",
        "email": "to",
        "send": "to",
        "kinh gui": "greeting_name",
        "cc": "cc",
        "bcc": "bcc",
        "noi dung": "body_main",
        "thong tin them": "extra_info",
        "tep dinh kem": "attachment_ref",
        "file dinh kem": "attachment_ref",
        "dinh kem": "attachment_ref",
        "attachment": "attachment_ref",
        "attachments": "attachment_ref",
    }

    REQUIRED_FIELDS = ("subject", "to", "body_main")

    def parse_excel(self, file_path: str) -> dict[str, Any]:
        try:
            from openpyxl import load_workbook
        except ModuleNotFoundError as exc:
            raise RuntimeError(
                "Thiếu package 'openpyxl'. Hãy cài openpyxl để dùng tính năng gửi email từ Excel."
            ) from exc

        path = Path(file_path)
        workbook = load_workbook(path, read_only=False, data_only=False)
        sheet = workbook.active

        header_row = next(sheet.iter_rows(min_row=1, max_row=1), None)
        if not header_row:
            workbook.close()
            raise ValueError("File Excel không có hàng tiêu đề.")

        column_map: dict[str, int] = {}
        for idx, header_cell in enumerate(header_row):
            key = self._map_header(header_cell.value)
            if key and key not in column_map:
                column_map[key] = idx

        missing = [field for field in self.REQUIRED_FIELDS if field not in column_map]
        if missing:
            workbook.close()
            raise ValueError(
                "Thiếu cột bắt buộc trong file Excel: "
                + ", ".join(sorted(missing))
            )

        valid_rows: list[BulkEmailRow] = []
        invalid_rows: list[dict[str, Any]] = []

        for row_number, row in enumerate(sheet.iter_rows(min_row=2), start=2):
            if not any((cell.value not in (None, "")) or cell.hyperlink for cell in row):
                continue

            values = {
                field: self._extract_cell_payload(row[idx], field) if idx < len(row) else ""
                for field, idx in column_map.items()
            }
            to = str(values.get("to", "")).lower()
            subject = str(values.get("subject", ""))
            body_main = str(values.get("body_main", ""))
            candidate_name = str(values.get("candidate_name", ""))
            greeting_name = str(values.get("greeting_name", "")) or candidate_name
            cc = self._normalize_email_list(str(values.get("cc", "")))
            bcc = self._normalize_email_list(str(values.get("bcc", "")))
            extra_info = str(values.get("extra_info", ""))
            attachment_refs = self._collect_attachment_refs(values)

            problems: list[str] = []
            if not to:
                problems.append("thiếu email")
            elif not EMAIL_RE.match(to):
                problems.append("email không hợp lệ")
            invalid_cc = self._find_invalid_emails(cc)
            if invalid_cc:
                problems.append("cc không hợp lệ: " + ", ".join(invalid_cc))
            invalid_bcc = self._find_invalid_emails(bcc)
            if invalid_bcc:
                problems.append("bcc không hợp lệ: " + ", ".join(invalid_bcc))
            if not subject:
                problems.append("thiếu tiêu đề")
            if not body_main:
                problems.append("thiếu nội dung")

            body_parts: list[str] = []
            if greeting_name:
                body_parts.append(f"Dear {greeting_name},")
            if body_main:
                body_parts.append(body_main)
            if extra_info:
                body_parts.append(extra_info)
            body = "\n\n".join(part.strip() for part in body_parts if part and part.strip()).strip()
            if not body:
                problems.append("không tạo được nội dung email")

            attachment_paths: list[str] = []
            if attachment_refs:
                attachment_paths, missing_attachments = self._resolve_attachment_paths(path, attachment_refs)
                if missing_attachments:
                    problems.append(
                        "không tìm thấy file đính kèm: " + ", ".join(missing_attachments)
                    )

            if problems:
                invalid_rows.append(
                    {
                        "row_number": row_number,
                        "email": to,
                        "candidate_name": candidate_name,
                        "cc": cc,
                        "bcc": bcc,
                        "attachment_paths": attachment_refs,
                        "reason": ", ".join(problems),
                    }
                )
                continue

            valid_rows.append(
                BulkEmailRow(
                    row_number=row_number,
                    to=to,
                    subject=subject,
                    body=body,
                    candidate_name=candidate_name,
                    greeting_name=greeting_name,
                    cc=cc,
                    bcc=bcc,
                    attachment_paths=attachment_paths,
                )
            )

        workbook.close()

        if not valid_rows:
            if invalid_rows:
                sample = " | ".join(
                    f"dòng {item.get('row_number')}: {item.get('reason')}"
                    for item in invalid_rows[:3]
                )
                raise ValueError(f"Không có dòng hợp lệ nào để gửi email. {sample}")
            raise ValueError("Không có dòng hợp lệ nào để gửi email.")

        return {
            "file_path": str(path),
            "sheet_name": sheet.title,
            "valid_rows": [asdict(item) for item in valid_rows],
            "invalid_rows": invalid_rows,
            "preview_rows": [asdict(item) for item in valid_rows],
            "total_valid": len(valid_rows),
            "total_invalid": len(invalid_rows),
        }

    def _extract_cell_payload(self, cell: Any, field: str) -> str:
        if not field.startswith("attachment_ref"):
            return _cell_text(getattr(cell, "value", cell))

        hyperlink = getattr(cell, "hyperlink", None)
        if hyperlink:
            target = getattr(hyperlink, "target", None) or getattr(hyperlink, "location", None)
            if target:
                normalized_target = self._normalize_attachment_ref(str(target))
                if normalized_target:
                    return normalized_target

        value = getattr(cell, "value", "")
        text = _cell_text(value)
        if text.startswith("=HYPERLINK("):
            parsed = self._parse_hyperlink_formula(text)
            if parsed:
                return self._normalize_attachment_ref(parsed)
        return self._normalize_attachment_ref(text)

    def _parse_hyperlink_formula(self, formula: str) -> str:
        match = re.match(r'=HYPERLINK\("([^"]+)"(?:\s*[;,]\s*"[^"]*")?\)', formula, re.IGNORECASE)
        if match:
            return match.group(1).strip()
        return ""

    def _normalize_attachment_ref(self, raw_ref: str) -> str:
        raw = (raw_ref or "").strip().strip("\"'")
        if not raw:
            return ""

        parsed = urlparse(raw)
        if parsed.scheme.lower() == "file":
            decoded_path = unquote(parsed.path or "")
            if re.match(r"^/[A-Za-z]:", decoded_path):
                decoded_path = decoded_path[1:]
            if parsed.netloc and not re.match(r"^[A-Za-z]:", decoded_path):
                decoded_path = f"//{parsed.netloc}{decoded_path}"
            return decoded_path.strip()

        return unquote(raw)

    def _map_header(self, raw_header: Any) -> str:
        normalized = _normalize_header(raw_header)
        if not normalized:
            return ""
        mapped = self.HEADER_MAP.get(normalized)
        if mapped:
            return mapped
        attachment_match = re.fullmatch(r"(?:tep|file) dinh kem (\d+)", normalized)
        if attachment_match:
            return f"attachment_ref_{attachment_match.group(1)}"
        return ""

    def _split_attachment_refs(self, attachment_ref: str) -> list[str]:
        raw = (attachment_ref or "").strip()
        if not raw:
            return []
        parts = re.split(r"[;\n]+", raw)
        return [part.strip().strip("\"'") for part in parts if part and part.strip()]

    def _collect_attachment_refs(self, values: dict[str, Any]) -> list[str]:
        refs: list[str] = []
        for field, value in values.items():
            if not field.startswith("attachment_ref"):
                continue
            refs.extend(self._split_attachment_refs(str(value)))
        return refs

    def _normalize_email_list(self, raw: str) -> str:
        emails = self._split_email_list(raw)
        return ", ".join(emails)

    def _split_email_list(self, raw: str) -> list[str]:
        text = (raw or "").strip()
        if not text:
            return []
        parts = re.split(r"[;,\n]+", text)
        return [part.strip().lower() for part in parts if part and part.strip()]

    def _find_invalid_emails(self, raw: str) -> list[str]:
        invalid: list[str] = []
        for email in self._split_email_list(raw):
            if not EMAIL_RE.match(email):
                invalid.append(email)
        return invalid

    def _resolve_attachment_paths(self, excel_path: Path, attachment_refs: list[str]) -> tuple[list[str], list[str]]:
        resolved_paths: list[str] = []
        missing_paths: list[str] = []

        for raw in attachment_refs:
            candidates: list[Path] = []
            candidate_path = Path(raw)
            if candidate_path.is_absolute():
                candidates.append(candidate_path)
            else:
                candidates.append((excel_path.parent / candidate_path).resolve())
                candidates.append(Path(os.path.expandvars(raw)).expanduser())

            resolved = ""
            for candidate in candidates:
                try:
                    candidate_resolved = candidate.resolve()
                except Exception:
                    continue
                if candidate_resolved.exists() and candidate_resolved.is_file():
                    resolved = str(candidate_resolved)
                    break

            if resolved:
                resolved_paths.append(resolved)
            else:
                missing_paths.append(raw)

        return resolved_paths, missing_paths

    def write_send_log(
        self,
        *,
        source_file: str,
        results: list[dict[str, Any]],
        total_success: int,
        total_failed: int,
    ) -> str:
        log_dir = ensure_runtime_dir("app_settings", "bulk_email_logs")
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        log_path = log_dir / f"bulk_email_{timestamp}.json"
        payload = {
            "source_file": source_file,
            "created_at": datetime.now().isoformat(timespec="seconds"),
            "total_success": total_success,
            "total_failed": total_failed,
            "results": results,
        }
        with log_path.open("w", encoding="utf-8") as f:
            json.dump(payload, f, ensure_ascii=False, indent=2)
        return str(log_path)

    def export_result_workbook(
        self,
        *,
        source_file: str,
        results: list[dict[str, Any]],
    ) -> str:
        try:
            from openpyxl import load_workbook
        except ModuleNotFoundError as exc:
            raise RuntimeError(
                "Thiếu package 'openpyxl'. Hãy cài openpyxl để xuất file Excel kết quả."
            ) from exc

        source_path = Path(source_file)
        workbook = load_workbook(source_path)
        sheet = workbook.active

        sent_col = sheet.max_column + 1
        failed_col = sheet.max_column + 2
        error_col = sheet.max_column + 3
        sheet.cell(row=1, column=sent_col, value="sent")
        sheet.cell(row=1, column=failed_col, value="failed")
        sheet.cell(row=1, column=error_col, value="error")

        result_by_row = {
            int(item["row_number"]): item
            for item in results
            if item.get("row_number")
        }

        for row_idx in range(2, sheet.max_row + 1):
            item = result_by_row.get(row_idx)
            if not item:
                continue
            status = str(item.get("status") or "").lower()
            error_text = str(item.get("error") or "")
            sheet.cell(row=row_idx, column=sent_col, value="yes" if status == "sent" else "")
            sheet.cell(row=row_idx, column=failed_col, value="yes" if status in {"failed", "skipped"} else "")
            sheet.cell(row=row_idx, column=error_col, value=error_text)

        export_dir = ensure_runtime_dir("app_settings", "bulk_email_logs")
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        output_path = export_dir / f"{source_path.stem}_result_{timestamp}.xlsx"
        workbook.save(output_path)
        workbook.close()
        return str(output_path)
